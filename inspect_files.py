import sys
import openpyxl
import docx
from pypdf import PdfReader

sys.stdout.reconfigure(encoding='utf-8')

print("--- EXCEL FILE INSPECTION ---")
wb = openpyxl.load_workbook("Phân nhóm thuyết trình Triết học Mác - Lênin.xlsx")
for sheet_name in wb.sheetnames:
    print(f"\nSheet: {sheet_name}")
    ws = wb[sheet_name]
    for r in range(1, min(ws.max_row + 1, 40)):
        row = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 12))]
        if any(row):
            print(f"R{r}: {row}")

print("\n--- DOCX FILE INSPECTION ---")
doc = docx.Document("mln111_598_cau_hoan_chinh.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

sample_p = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
print(f"Non-empty paragraphs: {len(sample_p)}")
print("First 30 non-empty paragraphs:")
for i, p in enumerate(sample_p[:30]):
    print(f"{i+1}: {p}")

print("\n--- PDF FILE INSPECTION ---")
reader = PdfReader("TRIẾT HỌC.pdf")
print(f"Total PDF pages: {len(reader.pages)}")
for p_num in range(min(15, len(reader.pages))):
    txt = reader.pages[p_num].extract_text()
    first_lines = txt.split('\n')[:3] if txt else ["(no text)"]
    print(f"Page {p_num+1}: {' | '.join(first_lines)}")
